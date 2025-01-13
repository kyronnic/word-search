from docx import Document
from docx.shared import Inches
from docx.enum.section import WD_SECTION
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Pt
from docx.oxml.table import CT_Row, CT_Tc
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

word_search_book = Document('test.docx')

sections = word_search_book.sections


def section_def(docx: Document()) -> Document():
    section = docx.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(5.83)
    section.page_length = Inches(8.27)
    # add in header?
    return docx


def table_format(table):
    tbl = table._tbl
    tcPr = tbl.tcPr
    tcBorders = OxmlElement("w:tcBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "nil")

    left = OxmlElement("w:left")
    left.set(qn("w:val"), "nil")

    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '0')

    right = OxmlElement("w:right")
    right.set(qn("w:val"), "nil")

    tcBorders.append(top)
    tcBorders.append(left)
    tcBorders.append(bottom)
    tcBorders.append(right)
    tcPr.append(tcBorders)


def print_word_search(board: list, docx: Document()):
    docx = section_def(docx)
    table = docx.add_table(rows=20, cols=20, style='Word Search')

    for i, row in enumerate(board):
        if len(board) < 20: #word list table
            word_list_section = docx.add_section(WD_SECTION.CONTINUOUS)
        else:
            table_row = table.rows[i].cells
            for j in range(20):
                cell = table_row[j]
                cell.text = row[j]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell.paragraphs[0].runs[0].font.size = Pt(8)
                cell.width = Inches(0.25)
    # table_format(table)
    docx.save('test.docx')


